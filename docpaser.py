#!/usr/bin/env python
# coding: utf-8

# In[1]:


# -*- coding: utf-8 -*-  
'''***********************************************************************
Author:Vivek
Date:2020/08/27

Description:
1. change doc to docx
2. parse docx
    a. parse table in docx
***********************************************************************'''
import os #用于获取目标文件所在路径
import configparser
from win32com import client as wc #导入模块
import docx
from docx import Document #导入库


'''Config information'''
cfg = configparser.ConfigParser()
cfg.read("config.ini",encoding='utf-8')
docx_path = cfg['DOCX']['filepath']
doc_path = cfg['DOC']['filepath']
docx_file = cfg['DOCX']['file']

'''
my color
'''
class mcolors:
    HEADER = '\033[95m'
    OKBLUE = '\033[94m'
    OKGREEN = '\033[92m'
    WARNING = '\033[93m'
    FAIL = '\033[91m'
    ENDC = '\033[0m'
    BOLD = '\033[1m'
    UNDERLINE = '\033[4m'

def doc2docx(doc_p,docx_p):
    doc_files=[]
    for file in os.listdir(doc_p):
        if file.endswith(".doc"): #排除文件夹内的其它干扰文件，只获取".doc"后缀的word文件
            doc_files.append(file) 
    
    word = wc.Dispatch("Word.Application") # 打开word应用程序
    for file in doc_files:
        print(file)
        doc = word.Documents.Open(doc_p+"/"+file) #打开word文件
        doc.SaveAs("{}x".format(docx_p+"/"+file), 12)#另存为后缀为".docx"的文件，其中参数12指docx文件
        doc.Close() #关闭原来word文件
    word.Quit()
    print("doc转换docx完成！")

def iter_unique_cells(row):
    """Generate cells in *row* skipping empty grid cells."""
    prior_tc = None
    for cell in row.cells:
        this_tc = cell._tc
        if this_tc is prior_tc:
            continue
        prior_tc = this_tc
        yield cell
def printcell(cell):
    prior_tc = cell._tc
    print("---------print cell------------")#每个table N x N个cell，CT_Tc为cell属性，合并的cell共用一个cell
    print("top: {} left: {}".format(prior_tc.top,prior_tc.left))#此cell的起点
    print("right: {} bottom: {}".format(prior_tc.right,prior_tc.bottom))
    print("_tr_idx: {} _grid_col: {}".format(prior_tc._tr_idx, prior_tc._grid_col))
    
def parsedocx_table_by_rows(table):
    print("**************************************")
    print("rows: {}".format(len(table.rows)))
    print("colums: {}".format(len(table.columns)))
    print("**************************************")
    for i, row in enumerate(table.rows[:]):   # 读每行
        row_content = []
        prior_tc = None
        for cell in row.cells:  # 读一行中的所有单元格
            this_tc = cell._tc
            if this_tc is prior_tc:
                c = " "
            else:   
                c = cell.text
                #printcell(cell)
                prior_tc = this_tc
            row_content.append(c)
        print ("{}. {}".format(i,row_content)) #以列表形式导出每一行数据
        i+=1
def parsedocx_table_by_cells(table):
    rows = len(table.rows)
    colums = len(table.columns)
    print("{}table-----------------------(◕ܫ◕){}".format(mcolors.OKBLUE,mcolors.ENDC))
    print("rows: {}".format(rows))
    print("colums: {}".format(colums))
    print("--------------------------------------------------------------------------")
    for i in range(rows):
        #row_content =[]
        row_content = ""
        for j in range(colums):
            cell = table.cell(i,j)
            prior_tc = cell._tc
            if i is prior_tc.top and j is prior_tc.left:
                c = cell.text
            else:
                c = ""
            #row_content.append("[{},{}] {}".format(i,j,c))
            row_content = row_content +" "+ mcolors.FAIL + "[" +str(i) +"," +str(j) +"]"+ mcolors.ENDC +c
        #print ("{}{}. {}{}".format(mcolors.OKGREEN,i,row_content,mcolors.ENDC)) #以列表形式导出每一行数据
        print("{}.{}".format(i,row_content))
    print("--------------------------------------------------------------------------")

def parsedocx(file):
    document = Document(file) #读入文件
    tables = document.tables #获取文件中的表格集
    for table in tables[:]:
        parsedocx_table_by_cells(table)
        #parsedocx_table_by_rows(table)

def main():
    parsedocx(docx_file)
    
if __name__ == "__main__":
    main()

